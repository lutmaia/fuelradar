"""Perfil exploratório local da ANP; não faz ingestão nem altera fontes.

Uso: python scripts/profile_anp_data.py "data/raw/*.xlsx"
"""

import argparse
from collections import Counter, defaultdict
import csv
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import glob
import hashlib
import io
from itertools import combinations
import json
from pathlib import Path
import re
import statistics
import sys
import tempfile
import unicodedata
from zipfile import ZipFile

from openpyxl import __version__ as openpyxl_version, load_workbook


class ProfileError(ValueError):
    """Arquivo ou esquema não interpretável com as regras conhecidas."""


def normalize_text(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())


UF_NAMES = dict(zip(
    "AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO".split(),
    "ACRE|ALAGOAS|AMAPA|AMAZONAS|BAHIA|CEARA|DISTRITO FEDERAL|ESPIRITO SANTO|GOIAS|MARANHAO|MATO GROSSO|MATO GROSSO DO SUL|MINAS GERAIS|PARA|PARAIBA|PARANA|PERNAMBUCO|PIAUI|RIO DE JANEIRO|RIO GRANDE DO NORTE|RIO GRANDE DO SUL|RONDONIA|RORAIMA|SANTA CATARINA|SAO PAULO|SERGIPE|TOCANTINS".split("|"),
))
UF_CODES = {name: code for code, name in UF_NAMES.items()}
PRODUCTS = {
    "GASOLINA": "GASOLINA_COMUM", "GASOLINA COMUM": "GASOLINA_COMUM",
    "GASOLINA ADITIVADA": "GASOLINA_ADITIVADA",
    "ETANOL": "ETANOL_HIDRATADO", "ETANOL HIDRATADO": "ETANOL_HIDRATADO",
    "DIESEL S500": "DIESEL_S500", "OLEO DIESEL": "DIESEL_S500",
    "DIESEL S10": "DIESEL_S10", "OLEO DIESEL S10": "DIESEL_S10",
    # A fonte histórica não especifica S500 neste rótulo: não fundir por suposição.
    "DIESEL": "DIESEL_NAO_ESPECIFICADO",
    "GNV": "GNV", "GLP": "GLP",
}
MISSING = {"", "-", "NA", "N/A", "NULL"}


def is_missing(value):
    return value is None or (isinstance(value, str) and normalize_text(value) in MISSING)


def normalize_uf(value):
    text = normalize_text(value)
    return text if text in UF_NAMES else UF_CODES.get(text)


def normalize_product(value):
    return PRODUCTS.get(normalize_text(value))


def normalize_unit(value):
    text = normalize_text(value).replace(" ", "")
    return {"R$/L": "BRL/L", "R$/LITRO": "BRL/L", "R$/M3": "BRL/M3",
            "R$/13KG": "BRL/13KG"}.get(text)


def parse_money(value):
    if is_missing(value):
        return None
    if isinstance(value, bool):
        raise ValueError("Booleano não é preço")
    text = str(value).strip()
    if "," in text:
        if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?:\.\d{3})+),\d+", text):
            raise ValueError(f"Valor monetário inválido: {value!r}")
        text = text.replace(".", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Valor monetário inválido: {value!r}") from exc
    if not result.is_finite():
        raise ValueError("Valor monetário não finito")
    return result


def parse_date(value):
    if is_missing(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Data inválida (esperado data Excel, DD/MM/AAAA ou ISO): {value!r}")


def normalize_identifier(value, width):
    """Retorna (texto, status). Não adivinha dígitos de texto curto/científico.

    Inteiros nativos Excel de até 14 dígitos têm precisão representável;
    completar sua largura fixa é reversível. Isso não valida o cadastro.
    """
    if is_missing(value):
        return None, "missing"
    if isinstance(value, bool):
        return None, "invalid"
    if isinstance(value, int):
        if 0 < value < 10 ** width:
            text = str(value)
            return text.zfill(width), "padded_native_integer" if len(text) < width else "exact_native_integer"
        return None, "invalid"
    if not isinstance(value, str):
        return None, "ambiguous_numeric"
    text = value.strip()
    mask = r"[0-9]{2}\.[0-9]{3}\.[0-9]{3}/[0-9]{4}-[0-9]{2}" if width == 14 else r"[0-9]{5}-[0-9]{3}"
    if re.fullmatch(mask, text):
        return re.sub(r"[./-]", "", text), "exact_masked_text"
    if re.fullmatch(r"[0-9]{" + str(width) + r"}", text):
        return text, "exact_text"
    if re.fullmatch(r"[0-9]{1," + str(width - 1) + r"}", text):
        return None, "ambiguous_short_text"
    if re.fullmatch(r"[0-9]+(?:[.,][0-9]+)?[eE][+-]?[0-9]+", text) or re.fullmatch(r"[0-9]+[.,][0-9]+", text):
        return None, "ambiguous_numeric_text"
    return None, "invalid"


def normalize_cnpj(value):
    return normalize_identifier(value, 14)[0]


ALIASES = {
    "UF": "uf", "ESTADO": "uf", "ESTADOS": "uf", "ESTADO - SIGLA": "uf",
    "MUNICIPIO": "municipality", "PRODUTO": "product", "BANDEIRA": "brand",
    "UNIDADE DE MEDIDA": "unit", "CNPJ": "cnpj", "CNPJ DA REVENDA": "cnpj",
    "DATA DA COLETA": "collection_date", "DATA DE COLETA": "collection_date",
    "PRECO DE REVENDA": "sale_price", "VALOR DE VENDA": "sale_price",
    "MES": "reference_month", "DATA INICIAL": "period_start", "DATA FINAL": "period_end",
    "REGIAO": "region", "BRASIL": "country", "PRECO MEDIO REVENDA": "mean_sale_price",
    "NUMERO DE POSTOS PESQUISADOS": "reported_station_count",
    "REGIAO - SIGLA": "region", "CEP": "postal_code",
    "REVENDA": "legal_name", "RAZAO": "legal_name", "FANTASIA": "trade_name",
    "NOME DA RUA": "address", "ENDERECO": "address",
    "NUMERO RUA": "address_number", "NUMERO": "address_number",
    "COMPLEMENTO": "address_complement", "BAIRRO": "neighborhood",
    "VALOR DE COMPRA": "purchase_price",
}


def identify_header(rows):
    """Consome até o cabeçalho e devolve linha, nomes, mapeamento e tipo."""
    for number, row in enumerate(rows, 1):
        if number > 50:
            break
        names = list(row)
        while names and names[-1] is None:
            names.pop()
        keys = [ALIASES.get(normalize_text(v)) for v in names]
        if "product" not in keys:
            continue
        if "cnpj" in keys or "sale_price" in keys or "collection_date" in keys:
            required = {"product", "uf", "municipality", "cnpj", "unit", "sale_price", "collection_date"}
            kind = "retail_observation"
        elif "reference_month" in keys:
            required = {"product", "unit", "mean_sale_price", "reference_month", "reported_station_count"}
            kind = "monthly_aggregate"
        elif "period_start" in keys:
            required = {"product", "unit", "mean_sale_price", "period_start", "period_end", "reported_station_count"}
            kind = "weekly_aggregate"
        else:
            raise ProfileError("Cabeçalho com PRODUTO, mas sem data/preço reconhecíveis")
        missing = required - set(keys)
        if missing:
            raise ProfileError(f"Cabeçalho na linha {number}: colunas obrigatórias ausentes: {sorted(missing)}")
        if kind != "retail_observation" and not set(keys) & {"uf", "region", "country", "municipality"}:
            raise ProfileError("Agregado sem dimensão geográfica reconhecida")
        if any(v is None for v in names) or len(set(map(normalize_text, names))) != len(names):
            raise ProfileError("Cabeçalho com coluna sem nome ou nome duplicado")
        return number, [str(v) for v in names], {k: i for i, k in enumerate(keys) if k}, kind
    raise ProfileError("Cabeçalho ANP não identificado nas primeiras 50 linhas")


def sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def date_range(values):
    return {"min": min(values).isoformat(), "max": max(values).isoformat()} if values else None


class Scope:
    """Acumula somente contagens, conjuntos de identidade e datas, sem linhas."""

    def __init__(self):
        self.rows = 0
        self.cities = set()
        self.stations = set()
        self.products = Counter()
        self.dates = set()
        self.automotive_rows = 0
        self.automotive_stations = set()
        self.automotive_cities = set()

    def add(self, city, station, product, dates):
        self.rows += 1
        if city:
            self.cities.add(city)
        if station:
            self.stations.add(station)
        self.products[product or "UNKNOWN"] += 1
        self.dates.update(d for d in dates if d)
        if product and product != "GLP":
            self.automotive_rows += 1
            if station:
                self.automotive_stations.add(station)
            if city:
                self.automotive_cities.add(city)

    def report(self, has_city=True, has_station=True):
        return {
            "rows": self.rows, "date_range": date_range(self.dates),
            "municipality_count": len(self.cities) if has_city else None,
            "municipalities": sorted(self.cities) if has_city else None,
            "unique_cnpj_count": len(self.stations) if has_station else None,
            "products": dict(sorted(self.products.items())),
            "automotive_rows": self.automotive_rows,
            "automotive_unique_cnpj_count": len(self.automotive_stations) if has_station else None,
            "automotive_municipality_count": len(self.automotive_cities) if has_city else None,
        }


def price_report(counts, values):
    result = dict(counts)
    if values:
        result.update(min=min(values), max=max(values), median=statistics.median(values))
        if len(values) >= 4:
            q1, _, q3 = statistics.quantiles(values, n=4, method="inclusive")
            low, high = q1 - 3 * (q3 - q1), q3 + 3 * (q3 - q1)
            result.update(iqr_fence=[low, high], potential_outliers=sum(v < low or v > high for v in values))
    return result


def profile_rows(rows, combined=None):
    rows = iter(rows)
    header_row, headers, mapping, kind = identify_header(rows)
    columns = [{"name": name, "types": Counter(), "missing": 0, "missing_markers": Counter(),
                "sp_missing": 0} for name in headers]
    dimensions = {k: Counter() for k in ("uf", "municipality", "product", "unit", "brand", "region") if k in mapping}
    dates = {k: set() for k in ("collection_date", "reference_month", "period_start", "period_end") if k in mapping}
    quality = Counter()
    duplicate_key_rows = []
    identifier_status = {k: Counter() for k in ("cnpj", "postal_code") if k in mapping}
    sp_identifier_status = {k: Counter() for k in identifier_status}
    ambiguous_identifier_rows = {k: [] for k in identifier_status}
    normalization_pairs = {k: Counter() for k in ("uf", "municipality", "product", "unit", "brand") if k in mapping}
    key_prices = {}
    all_hashes, sp_hashes, all_keys, sp_keys = set(), set(), set(), set()
    sp, jundiai = Scope(), Scope()
    price_columns = [i for i, h in enumerate(headers) if normalize_text(h).startswith("PRECO ") or normalize_text(h) in {"VALOR DE VENDA", "VALOR DE COMPRA"}]
    prices = defaultdict(lambda: (Counter(), []))
    sp_prices = defaultdict(lambda: (Counter(), []))
    total = blank = 0
    reported_counts = []
    for row_number, row in enumerate(rows, header_row + 1):
        if all(v is None or v == "" for v in row):
            blank += 1
            continue
        if len(row) > len(headers) and any(v is not None and v != "" for v in row[len(headers):]):
            raise ProfileError(f"Linha {row_number}: valores fora das colunas nomeadas")
        row = tuple(row[:len(headers)]) + (None,) * max(0, len(headers) - len(row))
        total += 1
        for stat, value in zip(columns, row):
            stat["types"][type(value).__name__] += 1
            if is_missing(value):
                stat["missing"] += 1
                stat["missing_markers"]["null" if value is None else str(value)] += 1
        get = lambda key: row[mapping[key]] if key in mapping else None
        for key in dimensions:
            if not is_missing(get(key)):
                dimensions[key][str(get(key))] += 1
        uf = normalize_uf(get("uf"))
        if "uf" in mapping and uf is None:
            quality["missing_or_unknown_uf"] += 1
        city = "" if is_missing(get("municipality")) else normalize_text(get("municipality"))
        if "municipality" in mapping and not city:
            quality["missing_municipality"] += 1
        product = normalize_product(get("product"))
        unit = normalize_unit(get("unit"))
        if product is None:
            quality["missing_or_unknown_product"] += 1
        if unit is None:
            quality["missing_or_unknown_unit"] += 1
        expected_unit = "BRL/13KG" if product == "GLP" else "BRL/M3" if product == "GNV" else "BRL/L"
        if product and unit and expected_unit != unit:
            quality["product_unit_mismatch"] += 1
        normalized_ids = {}
        for name, statuses in identifier_status.items():
            value, status = normalize_identifier(get(name), 14 if name == "cnpj" else 8)
            normalized_ids[name] = value
            statuses[status] += 1
            if uf == "SP":
                sp_identifier_status[name][status] += 1
            if status.startswith("ambiguous") and len(ambiguous_identifier_rows[name]) < 10:
                ambiguous_identifier_rows[name].append(row_number)
        station = normalized_ids.get("cnpj")
        if "cnpj" in mapping and station is None:
            quality["invalid_cnpj_format"] += 1
        normalized_dimensions = {"uf": uf, "municipality": city, "product": product, "unit": unit,
                                 "brand": None if is_missing(get("brand")) else normalize_text(get("brand"))}
        for name, counts in normalization_pairs.items():
            counts[(get(name), normalized_dimensions[name])] += 1
        parsed_dates = {}
        for key in dates:
            try:
                value = parse_date(get(key))
                if value:
                    dates[key].add(value)
                else:
                    quality[f"missing_{key}"] += 1
                parsed_dates[key] = value
            except ValueError:
                quality[f"invalid_{key}"] += 1
                parsed_dates[key] = None
        if parsed_dates.get("period_start") and parsed_dates.get("period_end"):
            if parsed_dates["period_start"] > parsed_dates["period_end"]:
                quality["inverted_period"] += 1
        digest = hashlib.sha256(json.dumps(row, default=str, ensure_ascii=False).encode()).digest()
        if digest in all_hashes:
            quality["exact_duplicate_rows"] += 1
        all_hashes.add(digest)
        geography = uf or normalize_text(get("region")) or normalize_text(get("country"))
        key = (geography, city, station, product, unit, *parsed_dates.values())
        key_complete = (product and unit and geography and all(parsed_dates.values())
                        and ("municipality" not in mapping or city)
                        and (station or kind != "retail_observation"))
        if key_complete:
            if key in all_keys:
                quality["candidate_key_duplicates"] += 1
                if len(duplicate_key_rows) < 10:
                    duplicate_key_rows.append(row_number)
            all_keys.add(key)
            if uf == "SP":
                if key in sp_keys:
                    quality["sp_candidate_key_duplicates"] += 1
                sp_keys.add(key)
        else:
            quality["rows_without_complete_candidate_key"] += 1
        if uf == "SP":
            for stat, value in zip(columns, row):
                if is_missing(value):
                    stat["sp_missing"] += 1
            if digest in sp_hashes:
                quality["sp_exact_duplicate_rows"] += 1
            sp_hashes.add(digest)
            args = (city, station, product, parsed_dates.values())
            sp.add(*args)
            if city == "JUNDIAI":
                jundiai.add(*args)
            if combined is not None and kind == "retail_observation":
                combined["sp"].add(*args)
                if city == "JUNDIAI":
                    combined["jundiai"].add(*args)
                if key_complete:
                    if key in combined["keys"]:
                        combined["duplicates"] += 1
                    combined["keys"].add(key)
            if kind == "retail_observation" and key_complete:
                try:
                    sale_price = parse_money(get("sale_price"))
                except ValueError:
                    sale_price = None
                values = key_prices.setdefault(key, set())
                if values and sale_price not in values:
                    quality["sp_candidate_key_price_conflicts"] += 1
                values.add(sale_price)
            if "reported_station_count" in mapping:
                try:
                    count = parse_money(get("reported_station_count"))
                    if count is not None:
                        reported_counts.append(float(count))
                except ValueError:
                    quality["invalid_reported_station_count"] += 1
        for index in price_columns:
            group = f"{headers[index]} | {product or normalize_text(get('product'))} | {unit or normalize_text(get('unit'))}"
            targets = [prices[group]] + ([sp_prices[group]] if uf == "SP" else [])
            for counts, values in targets:
                counts["rows"] += 1
                try:
                    value = parse_money(row[index])
                    if value is None:
                        counts["missing"] += 1
                    else:
                        values.append(float(value))
                        counts["numeric"] += 1
                        if value <= 0:
                            counts["negative" if value < 0 else "zero"] += 1
                except ValueError:
                    counts["non_numeric"] += 1
    if total == 0:
        raise ProfileError("Cabeçalho reconhecido, mas nenhuma linha de dados encontrada")
    for stat in columns:
        stat["missing_fraction"] = stat["missing"] / total if total else None
        stat["sp_missing_fraction"] = stat["sp_missing"] / sp.rows if sp.rows else None
    if combined is not None and kind == "retail_observation" and "sources" in combined:
        combined["sources"].append({"keys": key_prices, "dates": sp.dates.copy()})
    return {
        "kind": kind, "header_row": header_row, "column_count": len(headers), "rows": total,
        "blank_rows_after_header": blank, "columns": columns,
        "canonical_mapping": {k: headers[i] for k, i in mapping.items()},
        "dimensions": dimensions, "dates": {k: date_range(v) for k, v in dates.items()},
        "normalization_pairs": {k: [{"original": original, "normalized": normalized, "rows": n}
                                     for (original, normalized), n in counts.items()]
                                for k, counts in normalization_pairs.items()},
        "identifier_status": identifier_status, "sp_identifier_status": sp_identifier_status,
        "ambiguous_identifier_row_numbers": ambiguous_identifier_rows,
        "quality": quality, "candidate_key_duplicate_row_numbers": duplicate_key_rows,
        "sp_identifiable": "uf" in mapping,
        "sp": sp.report("municipality" in mapping, "cnpj" in mapping) if "uf" in mapping else None,
        "jundiai": jundiai.report(True, "cnpj" in mapping) if "municipality" in mapping and "uf" in mapping else None,
        "sp_reported_station_count_range": [min(reported_counts), max(reported_counts)] if reported_counts else None,
        "prices_all_source": {k: price_report(*v) for k, v in sorted(prices.items())},
        "prices_sp": {k: price_report(*v) for k, v in sorted(sp_prices.items())},
    }


def profile_xlsx(source, combined):
    with ZipFile(source) as archive:
        members = archive.namelist()
        encodings = set()
        for name in members:
            if name.endswith(".xml"):
                with archive.open(name) as stream:
                    prefix = stream.read(160)
                match = re.search(br'encoding=["\']([^"\']+)', prefix)
                encodings.add(match.group(1).decode() if match else "UTF-8 (XML default)")
    if hasattr(source, "seek"):
        source.seek(0)
    wb = load_workbook(source, read_only=True, data_only=True)
    sheets = []
    try:
        for ws in wb:
            dimensions = {"declared_rows": ws.max_row, "declared_columns": ws.max_column}
            # Não confiar em dimensões Excel, que podem conter somente formatação.
            ws.reset_dimensions()
            try:
                profile = profile_rows(ws.iter_rows(values_only=True), combined)
            except (ValueError, TypeError) as exc:
                raise ProfileError(f"Planilha {ws.title}: {exc}") from exc
            sheets.append({"sheet": ws.title, **dimensions, **profile})
    finally:
        wb.close()
    return {"format": "XLSX (ZIP/OOXML)", "encoding": sorted(encodings), "delimiter": None,
            "archive_members": members, "sheets": sheets}


def profile_csv(stream, combined):
    sample = stream.read(65536)
    stream.seek(0)
    try:
        decoded = sample.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        decoded = sample.decode("cp1252")
        encoding = "cp1252"
    try:
        dialect = csv.Sniffer().sniff(decoded, delimiters=";,\t|")
    except csv.Error as exc:
        raise ProfileError("Delimitador CSV não identificado") from exc
    wrapper = io.TextIOWrapper(stream, encoding=encoding, newline="")
    try:
        profile = profile_rows(csv.reader(wrapper, dialect), combined)
    finally:
        wrapper.detach()
    return {"format": "CSV", "encoding": encoding, "encoding_basis": "inferida dos primeiros 64 KiB; decodificação estrita do restante",
            "delimiter": dialect.delimiter, "sheets": [{"sheet": None, **profile}]}


def profile_file(path, combined):
    path = Path(path)
    before = sha256(path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        report = profile_xlsx(path, combined)
    elif suffix == ".csv":
        with path.open("rb") as stream:
            report = profile_csv(stream, combined)
    elif suffix == ".zip":
        with ZipFile(path) as archive:
            members = [m for m in archive.infolist() if not m.is_dir()]
            if not members or any(Path(m.filename).suffix.lower() not in {".csv", ".xlsx"} for m in members):
                raise ProfileError("ZIP com conteúdo inesperado: esperado somente CSV ou XLSX, sem ZIP aninhado")
            inner = []
            for member in members:
                # Arquivo temporário fora de raw; nunca extrair caminhos do ZIP.
                with tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024) as stream:
                    with archive.open(member) as source:
                        while chunk := source.read(1024 * 1024):
                            stream.write(chunk)
                    stream.seek(0)
                    item = profile_xlsx(stream, combined) if member.filename.lower().endswith(".xlsx") else profile_csv(stream, combined)
                    inner.append({"member": member.filename, **item})
            report = {"format": "ZIP", "archive_members": [m.filename for m in members], "contents": inner}
    else:
        raise ProfileError(f"Formato não suportado: {suffix}; use XLSX, CSV ou ZIP")
    if before != sha256(path):
        raise ProfileError(f"Arquivo mudou durante a leitura: {path.name}")
    return {"file": path.name, "size_bytes": path.stat().st_size, "sha256": before,
            "raw_unchanged": True, **report}


def tables(file):
    if "contents" in file:
        for item in file["contents"]:
            for sheet in item["sheets"]:
                yield f"{file['file']}!{item['member']}", sheet
    else:
        for sheet in file["sheets"]:
            yield file["file"], sheet


def summary_markdown(report):
    lines = ["# Perfil ANP — resultado gerado", "", "Contagens completas, sem amostragem de linhas. SP inclui GLP para diagnóstico; `automotive_*` exclui GLP.",
             "Datas de agregados são referências, não datas individuais de coleta. Postos únicos só são calculáveis com CNPJ.", "",
             "| Arquivo / planilha | Tipo | Linhas totais | Linhas SP | Período SP | Municípios SP | CNPJ únicos SP |", "|---|---|---:|---:|---|---:|---:|"]
    for file in report["files"]:
        for name, sheet in tables(file):
            sp = sheet["sp"] or {}
            period = sp.get("date_range") or {}
            lines.append(f"| {name} / {sheet['sheet'] or 'CSV'} | {sheet['kind']} | {sheet['rows']} | {sp.get('rows', 'N/D')} | {period.get('min', 'N/D')} a {period.get('max', 'N/D')} | {sp.get('municipality_count')} | {sp.get('unique_cnpj_count')} |")
    lines += ["", "Agregados em diferentes níveis não devem ser somados. `None`/N/D indica dimensão não disponível.", "", "## Esquemas distintos", ""]
    for group in report["schema_groups"]:
        lines += [f"- {group['kind']} ({len(group['columns'])} colunas, {len(group['sources'])} planilhas): " + "; ".join(group["columns"])]
    lines += ["", "## União das observações individuais SP", "", "```json", json.dumps(report["combined_retail"], ensure_ascii=False, indent=2), "```",
              "", "## Sobreposições entre fontes individuais (SP)", "", "```json", json.dumps(report.get("overlaps_sp", []), ensure_ascii=False, indent=2), "```",
              "", "Consulte profile.json para cabeçalhos, tipos, ausências, domínios originais/normalizados, preços, identificadores ambíguos, duplicidades e hashes."]
    return "\n".join(lines) + "\n"


def compare_sources(sources):
    """Compara somente chaves completas de observações SP; sem deduplicar."""
    reports = []
    for left, right in combinations(sources, 2):
        shared = left["keys"].keys() & right["keys"].keys()
        dates = left["dates"] & right["dates"]
        reports.append({
            "left": left["source"], "right": right["source"],
            "shared_collection_dates": len(dates), "shared_date_range": date_range(dates),
            "shared_candidate_keys": len(shared),
            "shared_keys_same_price_set": sum(left["keys"][k] == right["keys"][k] for k in shared),
            "shared_keys_conflicting_price_sets": sum(left["keys"][k] != right["keys"][k] for k in shared),
        })
    return reports


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", help="Arquivos ou padrões glob (entre aspas)")
    parser.add_argument("--output-dir", type=Path, default=Path("reports/data-profile"))
    args = parser.parse_args(argv)
    files = set()
    for pattern in args.paths:
        matches = glob.glob(pattern, recursive=True)
        if not matches:
            parser.error(f"Nenhum arquivo corresponde a {pattern}")
        files.update(Path(p).resolve() for p in matches)
    output = args.output_dir.resolve()
    raw = Path("data/raw").resolve()
    if output == raw or raw in output.parents or any(output == p.parent or p.parent in output.parents for p in files):
        parser.error("Diretório de saída deve ficar fora de raw e das pastas de origem")
    combined = {"sp": Scope(), "jundiai": Scope(), "keys": set(), "duplicates": 0, "sources": []}
    report = {"profile_version": 2, "generated_at_utc": datetime.now(timezone.utc).isoformat(),
              "python_version": sys.version.split()[0], "openpyxl_version": openpyxl_version,
              "script_sha256": sha256(__file__), "files": []}
    for path in sorted(files):
        source_start = len(combined["sources"])
        try:
            result = profile_file(path, combined)
        except Exception as exc:
            print(f"ERRO em {path.name}: {type(exc).__name__}: {exc}", file=sys.stderr)
            return 1
        report["files"].append(result)
        source_tables = [(name, sheet) for name, sheet in tables(result) if sheet["kind"] == "retail_observation"]
        for source, (name, sheet) in zip(combined["sources"][source_start:], source_tables, strict=True):
            source["source"] = {"input_path": path.as_posix(), "table": name, "sheet": sheet["sheet"]}
        print(f"OK {path.name}: {sum(t['rows'] for _, t in tables(result))} linhas em {sum(1 for _ in tables(result))} planilha(s); SHA-256 preservado", flush=True)
    report["combined_retail"] = {"sp": combined["sp"].report(), "jundiai": combined["jundiai"].report(),
                                 "candidate_key_duplicates_across_inputs": combined["duplicates"]}
    report["overlaps_sp"] = compare_sources(combined["sources"])
    groups = {}
    for file in report["files"]:
        for name, sheet in tables(file):
            signature = (sheet["kind"], tuple(c["name"] for c in sheet["columns"]))
            group = groups.setdefault(signature, {"kind": sheet["kind"], "columns": list(signature[1]), "sources": []})
            group["sources"].append({"file": name, "sheet": sheet["sheet"]})
    report["schema_groups"] = list(groups.values())
    output.mkdir(parents=True, exist_ok=True)
    (output / "profile.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (output / "summary.md").write_text(summary_markdown(report), encoding="utf-8")
    sp = report["combined_retail"]["sp"]
    print(f"SP (observações): {sp['rows']} linhas; {sp['municipality_count']} municípios; {sp['unique_cnpj_count']} CNPJ; {sp['date_range']}")
    print(f"Relatórios: {output / 'profile.json'} e summary.md")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    raise SystemExit(main())
