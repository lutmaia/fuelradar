"""Leitura das fontes externas para um formato único de linha.

Cada linha: `source_id`, `product` (código do projeto ou da fonte), `geo_level` (`pais`, `regiao`,
`uf`, `municipio`), `geo` (sigla ou nome normalizado), `uf`, `period_start`, `period_end`, `value`,
`stddev` e `unit`. Valor ausente na fonte (`-`, `***`) não vira linha: lacuna é lacuna.
"""

from datetime import date, datetime, timedelta
import json

from openpyxl import load_workbook

from ..ingestion.normalize import UF_CODES, normalize_text
from .sources import DISTRIBUTION, DISTRIBUTION_PRODUCTS, DOLLAR, PRODUCER_PRODUCTS, PRODUCER_REGIONS, PRODUCERS


class FormatError(Exception):
    """O arquivo não tem a estrutura esperada; nada dele é usado."""


def _row(source, product, level, geo, uf, start, end, value, unit, stddev=None):
    return {"source_id": source, "product": product, "geo_level": level, "geo": geo, "uf": uf,
            "period_start": start, "period_end": end, "value": round(float(value), 6),
            "stddev": None if stddev is None else round(float(stddev), 6), "unit": unit}


# -- Dólar (BCB, SGS 1) -----------------------------------------------------------------------


def parse_dollar(data):
    try:
        items = json.loads(data)
    except ValueError as exc:
        raise FormatError(f"Resposta do BCB não é JSON: {exc}") from exc
    if not isinstance(items, list) or not items:
        raise FormatError("Resposta do BCB vazia ou fora do formato de lista")
    rows = []
    for item in items:
        try:
            day = datetime.strptime(item["data"], "%d/%m/%Y").date()
            value = float(item["valor"])
        except (KeyError, TypeError, ValueError) as exc:
            raise FormatError(f"Linha do BCB fora do formato: {item!r}") from exc
        rows.append(_row(DOLLAR.source_id, "USD_BRL_VENDA", "pais", "BR", None, day, day, value, "BRL/USD"))
    return rows


# -- Produtores e importadores (ANP, .xls) ----------------------------------------------------


def _excel_date(value):
    return date(1899, 12, 30) + timedelta(days=int(value))


def parse_producers(path):
    import xlrd  # só esta fonte usa .xls
    try:
        sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
    except xlrd.XLRDError as exc:
        raise FormatError(f"Planilha de produtores ilegível: {exc}") from exc
    return parse_producers_sheet(sheet)


def parse_producers_sheet(sheet):
    """Lê uma planilha com a interface do xlrd (`nrows`, `ncols`, `cell_value`), testável sem arquivo."""
    header = next((r for r in range(min(sheet.nrows, 30))
                   if normalize_text(sheet.cell_value(r, 0)) == "PRODUTO"), None)
    if header is None or sheet.ncols < 9:
        raise FormatError("Cabeçalho 'Produto' não encontrado na planilha de produtores")
    regions = [normalize_text(sheet.cell_value(header + 1, c)) for c in range(3, 8)] + [
        normalize_text(sheet.cell_value(header, 8))]
    if tuple(regions) != PRODUCER_REGIONS:
        raise FormatError(f"Colunas de região inesperadas: {regions}")
    rows = []
    for r in range(header + 2, sheet.nrows):
        label, start, end = (sheet.cell_value(r, c) for c in range(3))
        product = PRODUCER_PRODUCTS.get(normalize_text(label))
        if product is None or not isinstance(start, float) or not isinstance(end, float):
            continue  # outros produtos, notas e linhas em branco
        for col, region in zip(range(3, 9), PRODUCER_REGIONS):
            value = sheet.cell_value(r, col)
            if isinstance(value, float):
                level = "pais" if region == "BRASIL" else "regiao"
                rows.append(_row(PRODUCERS.source_id, product, level, region, None, _excel_date(start),
                                 _excel_date(end), value, "BRL/L"))
    if not rows:
        raise FormatError("Nenhum preço dos produtos esperados na planilha de produtores")
    return rows


# -- Distribuição (ANP, .xlsx) ----------------------------------------------------------------

DISTRIBUTION_HEADER = ("DATA INICIAL", "DATA FINAL", "REGIAO", "ESTADO")


def parse_distribution(path, target_ufs):
    """Linhas de estado ou de município, só das UFs publicadas; o arquivo nacional fica intacto em raw."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        header, rows = None, []
        for values in wb.worksheets[0].iter_rows(values_only=True):
            if header is None:
                names = tuple(normalize_text(v) for v in values)
                if names[:4] == DISTRIBUTION_HEADER:
                    header = names
                    municipal = "MUNICIPIO" in header
                    need = {"PRODUTO", "UNIDADE DE MEDIDA", "PRECO MEDIO DISTRIBUICAO", "DESVIO PADRAO"}
                    if not need <= set(header):
                        raise FormatError(f"Colunas ausentes na distribuição: {sorted(need - set(header))}")
                continue
            rec = dict(zip(header, values))
            if not isinstance(rec["DATA INICIAL"], datetime):
                continue
            uf = UF_CODES.get(normalize_text(rec["ESTADO"]))
            if uf not in target_ufs:
                continue
            value = rec["PRECO MEDIO DISTRIBUICAO"]
            if not isinstance(value, (int, float)):
                continue  # "-": sem preço na semana
            product = DISTRIBUTION_PRODUCTS.get(normalize_text(rec["PRODUTO"]))
            if product is None:
                continue
            stddev = rec["DESVIO PADRAO"] if isinstance(rec["DESVIO PADRAO"], (int, float)) else None
            geo = normalize_text(rec["MUNICIPIO"]) if municipal else uf
            rows.append(_row(DISTRIBUTION.source_id, product, "municipio" if municipal else "uf", geo, uf,
                             rec["DATA INICIAL"].date(), rec["DATA FINAL"].date(), value,
                             normalize_text(rec["UNIDADE DE MEDIDA"]).replace("R$/L", "BRL/L"), stddev))
        if header is None:
            raise FormatError("Cabeçalho 'DATA INICIAL' não encontrado na planilha de distribuição")
        return rows
    finally:
        wb.close()
